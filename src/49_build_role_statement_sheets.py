# -*- coding: utf-8 -*-
"""
49_build_role_statement_sheets.py

Blind human coding of the role statement, across every condition that carries
relational evidence.

Why this round is needed
------------------------
The strict role-separation criterion is a regex, and it has now been shown to
split semantically identical sentences. These two say the same thing:

    "...provides contextual attribution, not primary risk prediction."      -> miss
    "...provides contextual attribution but is not the primary driver..."   -> hit

The first misses because "risk" sits between "primary" and "prediction"; the
pattern allows only an optional "stress " there. A second pattern misses
"Relational/network evidence" because it requires "relational" to be followed
immediately by context/attribution/evidence. On that basis the decoder-enforced
condition appeared to score 10/10 against 1/10 for the unenforced one, when
reading the texts shows both at 10/10.

Widening the regex against outputs already seen would be tuning the instrument
to the data. The alternative, and the one this project has used before, is to
let human coding be the primary measurement and to validate the regex against
it. That precedent is in 43_build_binary_rating_sheets.py: during calibration a
reader found "network density ... is not a primary predictor", the regex had
required "not the primary", and the script was corrected as a result.

What the readers code
---------------------
Two binary questions per note. They separate the two failure modes that the
single strict criterion collapses:

  S_CTX  Does the note describe relational/network evidence as context or
         attribution?  (the positive half)
  S_NOT  Does the note state explicitly that relational/network evidence is NOT
         the primary predictor/driver of the risk signal?  (the contrast)

The system instruction demands both. Llama 3 fills the required field with
"provides context" and never makes the contrast, so S_CTX and S_NOT come apart;
one number cannot show that.

Items
-----
Every successful generation that carries network evidence, across six blocks:
Llama 3 `full` (April main set), Llama 3 `full` (October, pinned weights),
Llama 3 `full_required_role`, GPT-OSS `full`, GPT-OSS `full_required_role`, and
GPT-OSS `full_required_role` under decoder enforcement. Conditions with no
network evidence are excluded: a note that makes no relational claim has no role
to mark, and scoring it as compliant is what produced the spurious 0.00-vs-1.00
contrast in the first place.

Blinding
--------
Model, condition, run and date are stripped. Each reader gets a different
random order. The key files are written separately and must not be opened
before both sheets are complete.

Output
------
outputs/llm_eval/role_statement_sheet_rater1.xlsx
outputs/llm_eval/role_statement_sheet_rater2.xlsx
outputs/llm_eval/role_statement_key_rater1.csv
outputs/llm_eval/role_statement_key_rater2.csv
"""

import importlib.util
import json
from pathlib import Path

import numpy as np
import pandas as pd

from config import PROJECT_ROOT
from common import ensure_dirs

EVAL = PROJECT_ROOT / "outputs" / "llm_eval"
GEN = PROJECT_ROOT / "outputs" / "llm_generations"
ensure_dirs(EVAL)

SRC = Path(__file__).resolve().parent


def load_render_note():
    """Reuse 43's renderer so readers see the same layout as the previous round."""
    path = SRC / "43_build_binary_rating_sheets.py"
    spec = importlib.util.spec_from_file_location("sheets43", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod.render_note


BLOCKS = [
    ("L-APR-FULL", "ollama_llama_native_10decisions_4conditions_final_clean.jsonl", "full"),
    ("L-OCT-FULL", "ollama_llama_full_rerun_q4_clean.jsonl", "full"),
    ("L-OCT-RR", "ollama_llama_required_role_q4.jsonl", "full_required_role"),
    ("G-FULL", "gptoss20b_schema.jsonl", "full"),
    ("G-RR", "gptoss20b_required_role.jsonl", "full_required_role"),
    ("G-RR-ENF", "gptoss20b_required_role_enforced.jsonl", "full_required_role"),
]

Q_CTX = "S_CTX: Not, ilişkisel/ağ kanıtını BAĞLAM ya da ATIF olarak niteliyor mu? (E/H)"
Q_NOT = "S_NOT: Not, ilişkisel/ağ kanıtının BİRİNCİL TAHMİN EDİCİ OLMADIĞINI açıkça söylüyor mu? (E/H)"
Q_NOTE = "Tereddüt notunuz (isteğe bağlı)"

DISPLAY = ["Sıra", "OKUYACAĞINIZ METİN", Q_CTX, Q_NOT, Q_NOTE]

INTRO = [
    ("NE YAPACAKSINIZ", ""),
    ("", "Aşağıda bir yapay zekâ modelinin ürettiği risk notları var. Her not için "
         "İKİ soruya E (evet) veya H (hayır) yazacaksınız. Başka bir şey "
         "değerlendirmeyeceksiniz: notun kalitesi, doğruluğu ya da yararlılığı "
         "sizi ilgilendirmiyor."),
    ("", ""),
    ("S_CTX nedir",
     "Not, ağ/ilişkisel kanıttan söz ederken bunu BAĞLAM, ATIF ya da eş anlamlı "
     "bir ifadeyle niteliyor mu? Örnek: 'ilişkisel bağlam sağlar', 'atıf amaçlıdır'. "
     "Sadece ağ verisini aktarıyorsa (yoğunluk 0.21, merkezî firma X) bu E DEĞİLDİR."),
    ("", ""),
    ("S_NOT nedir",
     "Not, ağ/ilişkisel kanıtın BİRİNCİL TAHMİN EDİCİ OLMADIĞINI açıkça söylüyor mu? "
     "Olumsuzlama şart. 'birincil tahmin edici değildir', 'riski asıl sürükleyen bu değil', "
     "'nedensellik kurmaz' gibi ifadeler E'dir. Sadece 'bağlam sağlar' demek, "
     "olumsuzlama içermiyorsa, H'dir."),
    ("", ""),
    ("ÖNEMLİ",
     "İfadenin hangi kelimelerle kurulduğu önemli DEĞİL. Anlam aynıysa E yazın. "
     "'birincil tahmin edici değildir' ile 'riskin asıl sürükleyicisi değildir' "
     "sizin için aynıdır."),
    ("", ""),
    ("", "Notların hangi koşuldan ya da hangi modelden geldiği kasıtlı olarak "
         "gizlenmiştir. Sıralama her okuyucu için farklıdır. Diğer okuyucuyla "
         "ikiniz de bitirmeden konuşmayın."),
]

EXAMPLE = {
    "Sıra": "ÖRNEK",
    "OKUYACAĞINIZ METİN":
        "İlişkisel ağ kanıtı: Ağ yoğunluğu 0.21, en merkezî firma MSFT. "
        "Bu kanıt ilişkisel bağlam sağlar, riskin birincil tahmin edicisi değildir.",
    Q_CTX: "E",
    Q_NOT: "E",
    Q_NOTE: "(örnek satır - doldurmayın)",
}

WIDTH = {"Sıra": 10, "OKUYACAĞINIZ METİN": 88, Q_CTX: 22, Q_NOT: 22, Q_NOTE: 30}


ROLE_HEADING = "■ RELATIONAL EVIDENCE ROLE STATEMENT"
NET_HEADING = "■ AĞ (İLİŞKİSEL) KANIT"
FIRM_HEADING = "■ FİRMA / SEKTÖR ATFI"


def defuse_blinding_leak(text):
    """Fold the dedicated role field into the relational-evidence section.

    The renderer gives every output field its own heading. Only the
    full_required_role condition has a relational_evidence_role_statement field,
    so that heading appears in exactly one condition and makes it identifiable
    on sight. Worse, the heading itself names the construct being coded, which
    primes the answer.

    Round 1 of the human evaluation failed partly for this reason: blinding
    leaked because the prompt condition is visible in the structure of the
    output. The sentence is therefore moved, unchanged, to the end of the
    relational-evidence section, where a note in any condition could have put
    it. No wording is altered; only the container is, exactly as the renderer
    already does for every other field.
    """
    if ROLE_HEADING not in text:
        return text

    blocks = text.split("\n\n")
    role_sentences, rest = [], []
    for b in blocks:
        if b.startswith(ROLE_HEADING):
            body = b[len(ROLE_HEADING):].strip()
            if body:
                role_sentences.append(body)
        else:
            rest.append(b)

    if not role_sentences:
        return "\n\n".join(rest)

    bullet = "• " + " ".join(role_sentences)

    for i, b in enumerate(rest):
        if b.startswith(NET_HEADING):
            rest[i] = b.rstrip() + "\n" + bullet
            return "\n\n".join(rest)

    for i, b in enumerate(rest):
        if b.startswith(FIRM_HEADING):
            rest.insert(i, NET_HEADING + "\n" + bullet)
            return "\n\n".join(rest)

    rest.append(NET_HEADING + "\n" + bullet)
    return "\n\n".join(rest)


def read_jsonl(p):
    return [json.loads(l) for l in open(p, encoding="utf-8") if l.strip()]


def has_network(rec):
    ids = rec.get("allowed_evidence_ids", []) or []
    return any(("NETWORK" in x) or x.startswith("FIRM-") or x.startswith("SECTOR-") for x in ids)


def collect():
    render_note = load_render_note()
    items = []
    for block, fname, cond in BLOCKS:
        path = GEN / fname
        if not path.exists():
            print(f"  ATLANDI (dosya yok): {fname}")
            continue
        kept = 0
        for rec in read_jsonl(path):
            if rec.get("status") != "ok":
                continue
            if rec.get("condition") != cond:
                continue
            if not has_network(rec):
                continue
            items.append({
                "block": block,
                "prompt_id": rec.get("prompt_id"),
                "decision_id": rec.get("decision_id"),
                "condition": cond,
                "model": rec.get("model"),
                "source_file": fname,
                "text": defuse_blinding_leak(render_note(rec.get("response_text", ""))),
            })
            kept += 1
        print(f"  {block:<12} {fname:<52} {kept} madde")
    return items


def build(items, seed, prefix):
    rng = np.random.default_rng(seed)
    order = rng.permutation(len(items))
    rows, key = [], []
    for pos, i in enumerate(order, start=1):
        it = items[i]
        item_id = f"{prefix}-{pos:03d}"
        rows.append({
            "Sıra": item_id,
            "OKUYACAĞINIZ METİN": it["text"],
            Q_CTX: "", Q_NOT: "", Q_NOTE: "",
        })
        key.append({"item_id": item_id, "prompt_id": it["prompt_id"],
                    "decision_id": it["decision_id"], "condition": it["condition"],
                    "model": it["model"], "block": it["block"],
                    "source_file": it["source_file"]})
    return pd.DataFrame(rows)[DISPLAY], pd.DataFrame(key)


def write_workbook(path, body):
    sheet = pd.concat([pd.DataFrame([EXAMPLE])[DISPLAY], body], ignore_index=True)
    try:
        with pd.ExcelWriter(path, engine="openpyxl") as xl:
            pd.DataFrame(INTRO, columns=["", " "]).to_excel(
                xl, sheet_name="ÖNCE BUNU OKUYUN", index=False)
            sheet.to_excel(xl, sheet_name="KODLAMA", index=False)

            ws0 = xl.sheets["ÖNCE BUNU OKUYUN"]
            ws0.column_dimensions["A"].width = 18
            ws0.column_dimensions["B"].width = 104
            for row in ws0.iter_rows():
                for c in row:
                    c.alignment = c.alignment.copy(wrapText=True, vertical="top")

            ws = xl.sheets["KODLAMA"]
            ws.freeze_panes = "C2"
            from openpyxl.utils import get_column_letter
            for idx, col in enumerate(DISPLAY, start=1):
                ws.column_dimensions[get_column_letter(idx)].width = WIDTH.get(col, 20)
            for row in ws.iter_rows():
                for c in row:
                    c.alignment = c.alignment.copy(wrapText=True, vertical="top")
        return True
    except PermissionError:
        print(f"  !! {path.name}: YAZILAMADI - dosya Excel'de acik veya OneDrive kilitli.")
        return False


def main():
    print("Maddeler toplaniyor:")
    items = collect()
    if not items:
        raise SystemExit("Hic madde bulunamadi.")
    print(f"\nToplam {len(items)} madde.\n")

    for rater, seed in [("rater1", 20261026), ("rater2", 20261027)]:
        body, key = build(items, seed, prefix="ROLE")
        ok = write_workbook(EVAL / f"role_statement_sheet_{rater}.xlsx", body)
        key.to_csv(EVAL / f"role_statement_key_{rater}.csv", index=False)
        if ok:
            print(f"  yazildi: role_statement_sheet_{rater}.xlsx  "
                  f"({len(body)} madde + 1 ornek satir)")
        print(f"  yazildi: role_statement_key_{rater}.csv")

    print("\nDagilim:")
    df = pd.DataFrame(items)
    print(df.groupby(["block", "model"]).size().to_string())
    print("\nAnahtar dosyalarini iki sayfa da bitmeden acmayin.")


if __name__ == "__main__":
    main()
